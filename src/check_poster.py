"""Find text that overflows its box, and boxes that collide, in a built poster.

    python src/check_poster.py outputs/thesis/EcoVision_poster_overview_A0.pptx

PowerPoint reflows text at open time and does not clip it. A textbox whose
content needs more room than the box simply draws over whatever is beneath,
which is invisible to python-pptx and to anything that only inspects shape
geometry. This measures each paragraph against the real font file, works out
the height the text actually needs, and reports:

  OVERFLOW   the text needs more vertical room than its box was given
  COLLISION  two shapes overlap once the overflow is taken into account

Run it after every layout change. It is much faster than opening the file,
and it catches the overlaps that a glance at a thumbnail misses.
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))
from build_poster import SANS, line_count  # noqa: E402

EMU_IN = 914400.0


def required_height(shape, factor: float) -> float:
    """Inches of vertical space this shape's text actually needs."""
    total = 0.0
    width_in = shape.width / EMU_IN / factor
    for p in shape.text_frame.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        text = "".join(r.text for r in runs)
        size = max((r.font.size.pt for r in runs if r.font.size), default=18)
        size /= factor
        name = next((r.font.name for r in runs if r.font.name), SANS)
        bold = any(r.font.bold for r in runs)
        spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
        n = line_count(text, width_in, size, font=name, bold=bold)
        total += n * size * spacing / 72.0
        if p.space_before is not None:
            total += p.space_before.pt / 72.0
        if p.space_after is not None:
            total += p.space_after.pt / 72.0
    return total * factor


def main() -> int:
    path = Path(sys.argv[1]) if len(sys.argv) > 1 else (
        REPO / "outputs/thesis/EcoVision_poster_overview_A0.pptx")
    prs = Presentation(path)
    slide = prs.slides[0]
    W = prs.slide_width / EMU_IN
    # The layout is authored at A0 and scaled; font sizes scale with it, so
    # measurement has to be done in the same space the text was authored in.
    factor = W / 33.11

    boxes = []
    problems = 0

    for shape in slide.shapes:
        if shape.left is None:
            continue
        geom = (shape.left / EMU_IN, shape.top / EMU_IN,
                shape.width / EMU_IN, shape.height / EMU_IN)
        need = geom[3]
        if shape.has_text_frame and shape.text_frame.text.strip():
            need = max(need, required_height(shape, factor))
            over = need - geom[3]
            if over > 0.12:
                problems += 1
                snippet = shape.text_frame.text.strip().replace("\n", " ")[:58]
                print(f"  OVERFLOW  +{over:5.2f} in  at x={geom[0]:5.2f} "
                      f"y={geom[1]:6.2f}  \"{snippet}\"")
        boxes.append((geom[0], geom[1], geom[0] + geom[2], geom[1] + need,
                      shape))

    # Collisions, ignoring shapes that are meant to sit on top of one
    # another: the panel fills, the accent strips and the numbered chips are
    # all deliberate backgrounds.
    #
    # Tables are content, not background. They have no text frame, so a
    # naive "no text means decoration" test skips them -- which is how a
    # five-row table sat across the footer while this script reported the
    # poster clean.
    def is_background(shape):
        if getattr(shape, "has_table", False) and shape.has_table:
            return False
        return not shape.has_text_frame or not shape.text_frame.text.strip()

    def describe(shape):
        if getattr(shape, "has_table", False) and shape.has_table:
            first = shape.table.cell(0, 0).text.strip()
            return f"[table: {first}]"
        return shape.text_frame.text.strip().replace("\n", " ")[:32]

    for i, (l1, t1, r1, b1, s1) in enumerate(boxes):
        for l2, t2, r2, b2, s2 in boxes[i + 1:]:
            if is_background(s1) or is_background(s2):
                continue
            overlap_x = min(r1, r2) - max(l1, l2)
            overlap_y = min(b1, b2) - max(t1, t2)
            if overlap_x > 0.3 and overlap_y > 0.18:
                problems += 1
                a, b = describe(s1), describe(s2)
                print(f"  COLLISION {overlap_y:5.2f} in  y={max(t1,t2):6.2f}  "
                      f"\"{a}\"  OVER  \"{b}\"")

    # Content that crosses a full-width separator rule. It may collide with
    # nothing and still look broken: a table hanging below the line that is
    # supposed to close the columns off reads as a mistake, and this is the
    # form the footer overflow took twice.
    rules = sorted(
        sh.top / EMU_IN for sh in slide.shapes
        if sh.left is not None
        and sh.width / EMU_IN > 0.8 * W
        and sh.height / EMU_IN < 0.12
        and sh.top / EMU_IN > 2.0
    )
    for l, t, r, b, shape in boxes:
        if (r - l) > 0.8 * W:
            continue
        for ry in rules:
            if t < ry - 0.3 < b and b > ry + 0.05:
                problems += 1
                print(f"  CROSSES RULE  {b - ry:5.2f} in below the rule at "
                      f"y={ry:6.2f}  x={l:5.2f}  \"{describe(shape)}\"")
                break

    print(f"\n{path.name}: {problems} problem(s)")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
